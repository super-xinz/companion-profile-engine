from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .template_people import TEMPLATE_BY_BIRTH_DATE, template_person_for_birth_date


SOURCE_FILES = {birth_date: person.source_file for birth_date, person in TEMPLATE_BY_BIRTH_DATE.items()}
SOURCE_META = {
    birth_date: {**person.birth_analysis, "mbti": person.mbti}
    for birth_date, person in TEMPLATE_BY_BIRTH_DATE.items()
}


def _source_path(birth_date: str) -> Path | None:
    person = template_person_for_birth_date(birth_date)
    if not person:
        return None
    filename = person.source_file
    project_root = Path(__file__).resolve().parents[2]
    candidates = [
        project_root / "source_profiles" / filename,
        Path.cwd() / "source_profiles" / filename,
        Path.cwd() / filename,
        Path.cwd().parent / filename,
    ]
    return next((path for path in candidates if path.exists()), None)


def _rows(sheet) -> list[list[Any]]:
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    while rows and not any(value is not None for value in rows[-1]):
        rows.pop()
    width = max((len(row) for row in rows), default=0)
    while width and all((row[width - 1] if len(row) >= width else None) is None for row in rows):
        width -= 1
    return [row[:width] for row in rows]


def load_source_document(birth_date: str) -> dict | None:
    path = _source_path(birth_date)
    if not path:
        return None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            "birth_date": birth_date,
            "source_file": path.name,
            "source_type": "user_supplied_complete_profile",
            "imported_at": datetime.now(timezone.utc).isoformat(),
            "sheets": {name: _rows(workbook[name]) for name in workbook.sheetnames},
        }
    finally:
        workbook.close()


def apply_source_profile(profile: dict, birth_date: str) -> bool:
    """Overlay the exact user-supplied profile and retain every original workbook cell."""
    document = load_source_document(birth_date)
    if not document:
        return False

    overview = document["sheets"].get("性格总览", [])
    meta = SOURCE_META[birth_date]
    profile["birth_analysis"].update({
        "bazi_text": meta["bazi_text"],
        "day_master": meta["day_master"],
        "pattern_name": meta["pattern_name"],
        "strength_label": meta["strength_label"],
        "relation_markers": meta["relation_markers"],
    })

    mbti_labels = {"E↔I": "ei", "S↔N": "sn", "T↔F": "tf", "J↔P": "jp"}
    imported_mbti = set()
    for row in overview:
        label = str(row[0] or "").replace(" ", "")
        key = next((value for prefix, value in mbti_labels.items() if label.startswith(prefix)), None)
        if not key or len(row) < 5:
            continue
        profile["mbti_dimensions"][key].update({
            "value": float(row[1]),
            "confidence": 0.45,
            "source_tendency": row[2],
            "source_interpretation": row[4],
        })
        imported_mbti.add(key)
    missing_mbti = set(mbti_labels.values()) - imported_mbti
    if missing_mbti:
        raise ValueError(f"{document['source_file']} 缺少MBTI维度: {sorted(missing_mbti)}")
    profile["mbti_dimensions"]["type_label"] = meta["mbti"]

    trait_lookup = {
        key: value
        for category in profile["core_traits"].values()
        for key, value in category.items()
    }
    imported_traits = set()
    for row in overview:
        if len(row) < 6:
            continue
        match = re.match(r"^([a-z_]+)", str(row[1] or ""))
        if not match or match.group(1) not in trait_lookup:
            continue
        key = match.group(1)
        tendency = str(row[4] or "").replace("⭐", "").strip()
        interpretation = str(row[5] or "").strip()
        trait_lookup[key].update({
            "value": float(row[2]),
            "confidence": 0.45,
            "tendency_label": tendency or trait_lookup[key]["tendency_label"],
            "interpretation": interpretation or tendency,
            "origin": "user_supplied_complete_profile",
        })
        imported_traits.add(key)
    missing_traits = set(trait_lookup) - imported_traits
    if missing_traits:
        raise ValueError(f"{document['source_file']} 缺少核心维度: {sorted(missing_traits)}")

    portrait_labels = {
        "本质": "essence",
        "优势": "strengths",
        "弱点": "weaknesses",
        "核心矛盾": "core_tension",
        "适合角色": "suitable_roles",
        "适合的角色": "suitable_roles",
    }
    source_portrait = {}
    for row in document["sheets"].get("人物画像", []):
        if len(row) < 2:
            continue
        key = portrait_labels.get(str(row[0] or "").strip())
        if key and row[1]:
            source_portrait[key] = {
                "content": row[1],
                "confidence": 0.45,
                "origin": "user_supplied_complete_profile",
            }
    profile["source_portrait"] = source_portrait
    missing_portrait = set(portrait_labels.values()) - set(source_portrait)
    if missing_portrait:
        raise ValueError(f"{document['source_file']} 缺少人物画像字段: {sorted(missing_portrait)}")
    profile["portrait"] = {
        key: {
            "content": item["content"],
            "confidence": 0.45,
            "origin": "user_supplied_complete_profile",
            "parameter_refs": [],
        }
        for key, item in source_portrait.items()
    }
    profile["source_profile_document"] = document
    profile["identity"]["template_person_id"] = template_person_for_birth_date(birth_date).user_id
    profile["meta"]["warnings"] = [
        warning for warning in profile["meta"].get("warnings", [])
        if "黄金样例" not in warning and "专家尚未提供任意生日到四位数字学编码" not in warning
    ]
    profile["meta"]["warnings"].append("已完整导入用户提供的原始画像工作簿；后续对话证据将持续校准动态画像。")
    return True
